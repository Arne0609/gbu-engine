# -*- coding: utf-8 -*-
"""Fragenkatalog der GBU-Variante „Riedl" mit der vollständigen Logik als xlsx.

    python3 gen_riedl_fragenkatalog_xlsx.py
    python3 gen_riedl_fragenkatalog_xlsx.py --out X.xlsx

Eingabe: norm_riedl_mf.json, norm_riedl_cyber.json, riedl_map.json
Ausgabe: GBU_Riedl_Fragenkatalog.xlsx

Zweck: Alles, was die Engine bei dieser Variante auswertet, in lesbarer Form –
Fragen mit Sichtbarkeit und Pflicht, Gefährdungen mit Anwendbarkeit, Regeln mit
Bedingung und Maßnahmen, dazu Karten, begründete Annahmen und die Vorbelegung
aus dem ZÜS-Prüfbericht. Bedingungen stehen als Klartext („[3.1 Aufzugsart] =
„Hydraulikaufzug" UND …"), Fragen werden mit ihrer Anzeigenummer genannt.

Blätter:
  Lesehinweise · Fragen · Gefährdungen · Regeln · Karten · Annahmen ·
  Nachweise · Maßnahmen
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
argv = sys.argv[1:]
OUT = os.path.join(HERE, 'GBU_Riedl_Fragenkatalog.xlsx')
if '--out' in argv:
    OUT = argv[argv.index('--out') + 1]
STAND = '20.09.2026'
FONT = 'Arial'


def lade(name):
    with open(os.path.join(HERE, name), encoding='utf-8') as fh:
        return json.load(fh)


MF = lade('norm_riedl_mf.json')
CY = lade('norm_riedl_cyber.json')
MAP = lade('riedl_map.json')

KATALOGE = [('GBU', MF), ('Cyber', CY)]

# Nachschlagewerke je Katalog: Fragen, Maßnahmen, Gefährdungen.
Q = {n: {q['code']: q for q in s['questions']} for n, s in KATALOGE}
M = {n: {m['code']: m for m in s['measures']} for n, s in KATALOGE}
H = {n: {h['code']: h for h in s['hazards']} for n, s in KATALOGE}

# Blatt und Zeile der Vorlage je Gefährdung (aus riedl_map.json).
BLATT = {}
for b in MAP['blaetter'] + [MAP['cyber']]:
    for z in b['zeilen']:
        for code in z['hazards']:
            BLATT[code] = (b, z)

STUFE = {'HIGH': 'Hoch', 'MEDIUM': 'Mittel', 'LOW': 'Niedrig',
         'NO_RISK': 'Kein Risiko', 'NOT_APPLICABLE': 'nicht zutreffend',
         'INCOMPLETE': 'unvollständig'}
ROLLE = {'TRIGGER': 'Befund', 'APPLICABILITY': 'Anwendbarkeit',
         'COMPENSATION': 'Kompensation', 'MODIFIER': 'Modifikator',
         'DOCUMENTATION': 'Dokumentation', 'OPTIONAL': 'optional'}
TYP = {'TECHNICAL': 'technisch', 'ORGANISATIONAL': 'organisatorisch'}
FREIGABE = {'VERIFIED': 'freigegeben', 'REVIEW_REQUIRED': 'Freigabe offen'}
HERKUNFT = {'RECONSTRUCTED_ORIGINAL': 'aus GBU 3.0 rekonstruiert',
            'NORM_DERIVED': 'aus der Norm abgeleitet', 'OWN_RULE': 'eigene Regel'}
EVIDENZ = {'HIGH_CONFIDENCE': 'belegt', 'DIRECT': 'direkt', 'INFERRED': 'abgeleitet',
           'HYPOTHESIS': 'Annahme'}
PHASE = {'stamm': 'Anlagenstamm', 'vorab': 'vorab (Betreiber)', 'vor_ort': 'vor Ort'}


# ------------------------------------------------------------------ Texte
def wertText(kat, qc, v):
    """Antwortwert lesbar: Optionslabel, Ja/Nein, Zahl mit Komma."""
    q = Q[kat].get(qc)
    if q and q.get('options'):
        for o in q['options']:
            if o['value'] == v:
                return '„%s"' % o['label']
    if v is True:
        return 'Ja'
    if v is False:
        return 'Nein'
    if isinstance(v, float):
        return str(v).replace('.', ',')
    return str(v)


# Prüfbericht 20.09.2026 (Befund zu Nummernkollisionen): GBU und Cyber führen
# eigene Nummernkreise, die sich überschneiden – [4.5] ist im GBU-Teil der
# UCM-Schutz, im Cyber-Teil der Softwarestand. In DIESER Datei stehen beide
# Kataloge nebeneinander, deshalb tragen Cyber-Nummern hier das Präfix „C".
# In der App und im Seed bleibt die Nummer unverändert (4.5).
NR_PRAEFIX = {'GBU': '', 'Cyber': 'C '}


def frageName(kat, qc, kurz=False):
    """Frage als „3.1 Aufzugsart …" oder kurz als „3.1" (Cyber: „C 3.1")."""
    q = Q[kat].get(qc)
    if not q:
        return qc
    nr = q.get('ui_number')
    nr = (NR_PRAEFIX.get(kat, '') + nr) if nr else qc
    return nr if kurz else '%s %s' % (nr, q['text'])


def ausdruck(kat, e, tiefe=0, kurz=True):
    """Bedingung der Engine als deutscher Satz."""
    if not e:
        return ''
    if 'all' in e:
        teile = [ausdruck(kat, x, tiefe + 1, kurz) for x in e['all']]
        if tiefe == 0:
            return '\nUND '.join(teile)
        return '(' + ' UND '.join(teile) + ')'
    if 'any' in e:
        teile = [ausdruck(kat, x, tiefe + 1, kurz) for x in e['any']]
        s = ' ODER '.join(teile)
        return s if tiefe == 0 else '(' + s + ')'
    if 'not' in e:
        return 'NICHT ' + ausdruck(kat, e['not'], tiefe + 1, kurz)
    qc, op, v = e.get('question'), e.get('operator'), e.get('value')
    n = frageName(kat, qc, kurz)
    if op == 'ANSWERED':
        return '[%s] ist beantwortet' % n
    if op == 'NOT_ANSWERED':
        return '[%s] ist unbeantwortet' % n
    ops = {'EQ': '=', 'NEQ': '≠', 'GT': '>', 'GTE': '≥', 'LT': '<', 'LTE': '≤',
           'IN': 'ist eines von', 'NOT_IN': 'ist keines von'}
    if isinstance(v, list):
        return '[%s] %s %s' % (n, ops.get(op, op),
                               ' / '.join(wertText(kat, qc, x) for x in v))
    return '[%s] %s %s' % (n, ops.get(op, op), wertText(kat, qc, v))


def anwendbarkeit(kat, h):
    """Wann die Gefährdung überhaupt bewertet wird."""
    teile = []
    appl = [x for x in h.get('questions', []) if x['role'] == 'APPLICABILITY']
    for x in appl:
        if x.get('applicable_when'):
            teile.append(ausdruck(kat, x['applicable_when'], 1))
        elif not any(y.get('applicable_when') for y in appl):
            teile.append('[%s] = Ja' % frageName(kat, x['question']))
    return ' UND '.join(dict.fromkeys(teile)) if teile else 'immer'


def pflichtfragen(kat, h):
    teile = []
    for x in h.get('questions', []):
        mode = x.get('required_mode', 'NEVER')
        if mode == 'ALWAYS':
            teile.append(frageName(kat, x['question'], True))
        elif mode == 'CONDITIONAL':
            teile.append('%s (wenn %s)' % (frageName(kat, x['question'], True),
                                           ausdruck(kat, x.get('required_when'), 1)))
    return ', '.join(teile) or '–'


def istAuffangregel(r):
    """Auffangregel „Kein Risiko" (mf_content/common.py, _catch_all).

    Erzeugt wird sie mit Priorität 1, Ergebnis „Kein Risiko" und der Bedingung
    „Ankerfrage ist beantwortet". Wörtlich gelesen („[5.1] ist beantwortet")
    ist das irreführend: Sie greift nur, wenn KEINE Mangelregel zutrifft – die
    Priorität 1 ist die eigentliche Bedingung. Deshalb wird sie im Bericht als
    Satz ausgeschrieben (Prüfbericht 20.09.2026, Befund zur Lesbarkeit).
    """
    c = r.get('condition')
    return (r.get('result') == 'NO_RISK' and r.get('priority') == 1
            and isinstance(c, dict) and c.get('operator') == 'ANSWERED')


def regelBedingung(kat, r):
    if istAuffangregel(r):
        return ('Auffangregel: keine Mangelregel dieser Gefährdung trifft zu und '
                'alle Pflichtfragen sind beantwortet (Anker: [%s]).'
                % frageName(kat, r['condition']['question'], True))
    return ausdruck(kat, r.get('condition'), 0)


def massnahmen(kat, r, gruppe):
    out = []
    for mm in r.get('measures', []):
        if (mm.get('group_id') == 'sofort') != (gruppe == 'sofort'):
            continue
        m = M[kat].get(mm['measure'])
        if m:
            out.append('• %s [%s]' % (m.get('title', ''), TYP.get(m.get('type'), '?')))
    return '\n'.join(out)


# ------------------------------------------------------------------ Layout
HEAD_FILL = PatternFill('solid', fgColor='1F3A5F')
ZWISCHEN_FILL = PatternFill('solid', fgColor='DCE6F1')
STUFE_FILL = {'HIGH': 'F4CCCC', 'MEDIUM': 'FCE5CD', 'LOW': 'FFF2CC',
              'NO_RISK': 'D9EAD3'}
thin = Side(style='thin', color='C9D1DB')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')


def blatt(wb, titel, spalten, breiten):
    ws = wb.create_sheet(titel)
    ws.append(spalten)
    for i, breite in enumerate(breiten, start=1):
        ws.column_dimensions[get_column_letter(i)].width = breite
    for zelle in ws[1]:
        zelle.font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
        zelle.fill = HEAD_FILL
        zelle.alignment = Alignment(wrap_text=True, vertical='center')
        zelle.border = BORDER
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 30
    return ws


def zeile(ws, werte, breiten, fuellung=None):
    ws.append(werte)
    nr = ws.max_row
    hoehe = 1
    for i, wert in enumerate(werte):
        text = '' if wert is None else str(wert)
        umbrueche = text.count('\n') + 1
        laenge = max((len(t) for t in text.split('\n')), default=0)
        hoehe = max(hoehe, umbrueche + int(laenge / max(breiten[i], 1)))
        z = ws.cell(row=nr, column=i + 1)
        z.font = Font(name=FONT, size=9)
        z.alignment = WRAP
        z.border = BORDER
        if fuellung:
            z.fill = fuellung
    ws.row_dimensions[nr].height = min(15 * min(hoehe, 12), 180)
    return nr


def abschluss(ws, spaltenzahl):
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(spaltenzahl), ws.max_row)


# ------------------------------------------------------------------ Blätter
wb = Workbook()
wb.remove(wb.active)

# --- Lesehinweise -----------------------------------------------------------
ws = wb.create_sheet('Lesehinweise')
ws.column_dimensions['A'].width = 34
ws.column_dimensions['B'].width = 118
kopfzeilen = [
    ('GBU-Variante „Riedl" – Fragenkatalog mit Logik', ''),
    ('Stand', STAND),
    ('Regelwerk GBU', '%s – %d Fragen, %d Gefährdungen, %d Regeln, %d Maßnahmen'
     % (MF['rule_version'], len(MF['questions']), len(MF['hazards']),
        len(MF['rules']), len(MF['measures']))),
    ('Regelwerk Cyber', '%s – %d Fragen, %d Gefährdungen, %d Regeln, %d Maßnahmen'
     % (CY['rule_version'], len(CY['questions']), len(CY['hazards']),
        len(CY['rules']), len(CY['measures']))),
    ('Herkunft', 'Abgeleitet aus dem mehrfragigen Katalog DIN EN 81-20 und dem '
     'Cyber-Katalog: gleiche Fragen, Gefährdungen, Regeln und Codes, nur der '
     'Umfang ist kleiner. Erzeugt von engine_model/gen_riedl_fragenkatalog_xlsx.py '
     'aus norm_riedl_mf.json, norm_riedl_cyber.json und riedl_map.json.'),
    ('', ''),
    ('So liest sich die Logik', ''),
    ('Fragenummer', 'Bedingungen nennen Fragen mit ihrer Anzeigenummer in eckigen '
     'Klammern, z. B. [3.1] = „Hydraulikaufzug". Das Blatt „Fragen" führt Nummer, '
     'Code und Text zusammen. GBU und Cyber haben eigene, sich überschneidende '
     'Nummernkreise ([4.5] ist im GBU-Teil der UCM-Schutz, im Cyber-Teil der '
     'Softwarestand). In dieser Datei tragen Cyber-Nummern deshalb das Präfix „C" '
     '([C 4.5]); in der App heißt die Frage unverändert 4.5.'),
    ('Rolle einer Frage', 'Befund = die Frage löst die Mangelregel aus. Anwendbarkeit = '
     'sie entscheidet, ob die Gefährdung überhaupt bewertet wird. Kompensation = sie '
     'kann die Stufe senken. Modifikator = sie verändert die Stufe, ohne sie allein '
     'auszulösen (Nutzerkreis, Ausführungsmerkmal). Dokumentation = ohne Stufenwirkung. '
     'Die Rolle beschreibt die Absicht; welche Regel tatsächlich greift, steht im Blatt '
     '„Regeln".'),
    ('Unauffälliger Wert', 'Nur gesetzt, wo eine Sammelantwort zulässig ist – also dort, '
     'wo der unauffällige Zustand mit einem Blick festzustellen ist. Messwerte, '
     'Bauart-, Ausstattungs- und Prüfaussagen tragen bewusst keinen Wert und werden '
     'einzeln beantwortet; eine leere Zelle ist deshalb kein Fehler.'),
    ('Sichtbar wenn', 'Die Frage erscheint im Fragebogen nur, wenn die Bedingung '
     'zutrifft. Ohne Eintrag ist sie immer sichtbar.'),
    ('Pflicht', 'Pflichtfragen müssen beantwortet sein, sonst bleibt die Gefährdung '
     '„unvollständig" (fail-closed) – sie wird nie stillschweigend grün.'),
    ('Anwendbar wenn', 'Trifft die Bedingung nicht zu, ist die Gefährdung „nicht '
     'zutreffend" und zählt in keiner Ampel mit.'),
    ('Regeln', 'Je Gefährdung werden alle Regeln geprüft. Es gewinnt die zutreffende '
     'Regel mit der höchsten Priorität; bei den Aggregationsarten MAXIMUM und ANY '
     'gelten zusätzlich die Maßnahmen weiterer zutreffender Befundregeln. Die '
     'Auffangregel „Kein Risiko" greift, wenn keine Mangelregel zutrifft und alle '
     'Pflichtfragen beantwortet sind.'),
    ('Stufe → Ampel (Riedl)', 'Kein Risiko und Niedrig = Grün, Mittel = Gelb, Hoch = '
     'Rot. Eine Zeile des Berichts zeigt die schlechteste Farbe ihrer Gefährdungen; '
     'eine unvollständige hält die Zeile offen.'),
    ('Freigabe', 'VERIFIED = fachlich freigegeben, REVIEW_REQUIRED = Freigabe steht '
     'aus. Die Variante erbt den Stand der Quellkataloge: GBU %d von %d Regeln '
     'freigegeben, Cyber %d von %d. Inhaltlich geänderte Regeln fallen dabei '
     'zurück auf „Freigabe offen" – das ist der Stand nach der Überarbeitung vom '
     '%s.'
     % (sum(1 for r in MF['rules'] if r.get('quality_status') == 'VERIFIED'),
        len(MF['rules']),
        sum(1 for r in CY['rules'] if r.get('quality_status') == 'VERIFIED'),
        len(CY['rules']), STAND)),
    ('Annahmen', 'Aus dem Baujahr begründete Vorbelegungen. Sie werden nicht '
     'gespeichert, sondern bei jeder Bewertung neu abgeleitet, und entfallen, sobald '
     'die Frage von Hand beantwortet wird oder der Konformitätsnachweis fehlt.'),
    ('Nachweise', 'Vorbelegung aus dem letzten ZÜS-Prüfbericht (TRBS 1201 Teil 4, '
     '3.3 (2)). Sie wird gespeichert, ist an das Datum des Berichts gebunden und '
     'wird vor Ort bestätigt.'),
    ('Karten', 'Fragengruppen: mehrere Einzelfragen als eine Karte. Auf einer '
     'Ankreuzliste ist ein Häkchen die Feststellung; „unauffällig" entsteht erst '
     'durch das Bestätigen der Karte. Der Cyber-Teil nutzt ausschließlich '
     'gemischte Karten mit unveränderten Fragen: Cyber-Befunde sind keine '
     'Sichtprüfung, ein Sammel-„unauffällig" würde dort unbelegte Tatsachen '
     'behaupten.'),
    ('Auffangregel', 'Die Regel mit Priorität 1 und Ergebnis „Kein Risiko" ist die '
     'Auffangregel der Gefährdung. Sie greift, wenn keine Mangelregel zutrifft und '
     'alle Pflichtfragen beantwortet sind; die genannte Ankerfrage stellt nur '
     'sicher, dass überhaupt erhoben wurde.'),
]
for links, rechts in kopfzeilen:
    ws.append([links, rechts])
    nr = ws.max_row
    a, b = ws.cell(row=nr, column=1), ws.cell(row=nr, column=2)
    a.font = Font(name=FONT, bold=True, size=11 if nr == 1 else 10)
    b.font = Font(name=FONT, size=10)
    a.alignment = WRAP
    b.alignment = WRAP
    ws.row_dimensions[nr].height = max(15, 13 * (int(len(str(rechts)) / 110) + 1))

# --- Fragen -----------------------------------------------------------------
SP_F = ['Katalog', 'Nr.', 'Code', 'Erhebungsbereich', 'Phase', 'Karte', 'Frage',
        'Typ', 'Antwortmöglichkeiten', 'Sichtbar wenn', 'Pflicht für',
        'Unauffälliger Wert', 'Aus dem Anlagenstamm', 'Hinweistext',
        'Bewertet in Gefährdungen']
BR_F = [7, 7, 22, 26, 13, 20, 52, 10, 34, 40, 34, 14, 18, 44, 26]
ws = blatt(wb, 'Fragen', SP_F, BR_F)
for kat, seed in KATALOGE:
    gruppen = {g['id']: g for g in seed.get('question_groups', [])}
    phasen = seed.get('category_phases') or {}
    # Je Frage: in welchen Gefährdungen kommt sie mit welcher Rolle vor?
    benutzt, pflicht = {}, {}
    for h in seed['hazards']:
        for x in h.get('questions', []):
            benutzt.setdefault(x['question'], []).append(
                '%s (%s)' % (h['code'], ROLLE.get(x['role'], x['role'])))
            mode = x.get('required_mode', 'NEVER')
            if mode == 'ALWAYS':
                pflicht.setdefault(x['question'], []).append('%s immer' % h['code'])
            elif mode == 'CONDITIONAL':
                pflicht.setdefault(x['question'], []).append(
                    '%s wenn %s' % (h['code'], ausdruck(kat, x.get('required_when'), 1)))
    for q in seed['questions']:
        g = gruppen.get(q.get('group'))
        karte = '%s %s' % (g.get('ui_number', ''), g['title']) if g else ''
        optionen = '\n'.join('%s = „%s"' % (o['value'], o['label'])
                             for o in q.get('options', []))
        if q['type'] == 'NUMBER' and (q.get('min') is not None or q.get('max') is not None):
            optionen = 'Zahl von %s bis %s' % (q.get('min', '–'), q.get('max', '–'))
        bc = q.get('best_case')
        zeile(ws, [
            kat,
            q.get('ui_number', ''),
            q['code'],
            q.get('category', ''),
            PHASE.get(phasen.get(q.get('category', '')), ''),
            karte,
            q['text'] + (' [optional]' if q.get('optional') else ''),
            q['type'],
            optionen,
            ausdruck(kat, q.get('visible_when'), 0),
            '\n'.join(pflicht.get(q['code'], [])) or '–',
            '' if bc is None else wertText(kat, q['code'], bc),
            q.get('stamm_key', ''),
            q.get('help_text', ''),
            ', '.join(benutzt.get(q['code'], [])) or '–',
        ], BR_F)
abschluss(ws, len(SP_F))

# --- Gefährdungen -----------------------------------------------------------
SP_G = ['Katalog', 'Code', 'Gefährdung', 'Erhebungsbereich', 'Baugruppe',
        'Blatt der Vorlage', 'Zeile', 'TRBS 3121 Anh. 1', 'Anwendbar wenn',
        'Pflichtfragen', 'Alle Fragen (Rolle)', 'Auswertung', 'Gefährdungsfaktor',
        'Betroffene', 'Quellen', 'Regeln']
BR_G = [7, 12, 40, 26, 22, 20, 26, 14, 40, 40, 40, 22, 20, 20, 28, 8]
ws = blatt(wb, 'Gefährdungen', SP_G, BR_G)
for kat, seed in KATALOGE:
    regelzahl = {}
    for r in seed['rules']:
        regelzahl[r['hazard']] = regelzahl.get(r['hazard'], 0) + 1
    for h in seed['hazards']:
        b, z = BLATT.get(h['code'], (None, None))
        zeile(ws, [
            kat,
            h['code'],
            h['title'],
            # Bei Gefährdungen ist ui_group der Erhebungsbereich und category
            # die Baugruppe (so gruppiert auch die App: MfKatalog.baugruppen).
            h.get('ui_group', ''),
            h.get('category', ''),
            '%s %s' % (b['kuerzel'], b['titel']) if b else '–',
            '%s%s %s' % (b['kuerzel'], z['nr'], z['titel']) if z else '–',
            ', '.join(str(n) for n in (z.get('trbs_anhang1') or [])) if z else '',
            anwendbarkeit(kat, h),
            pflichtfragen(kat, h),
            '\n'.join('%s – %s' % (frageName(kat, x['question'], True),
                                   ROLLE.get(x['role'], x['role']))
                      for x in h.get('questions', [])),
            '%s / %s' % (h.get('aggregation_type', ''), h.get('evaluation_mode', '')),
            h.get('hazard_factor', ''),
            ', '.join(h.get('person_groups', [])),
            '\n'.join(('%s %s' % (s.get('document', ''), s.get('section', ''))).strip()
                      for s in h.get('sources', [])),
            regelzahl.get(h['code'], 0),
        ], BR_G)
abschluss(ws, len(SP_G))

# --- Regeln -----------------------------------------------------------------
SP_R = ['Katalog', 'Regel', 'Gefährdung', 'Titel der Gefährdung', 'Priorität',
        'Wenn', 'Nur anwendbar wenn', 'Ergebnis', 'Sofortmaßnahmen',
        'Mittelfristige Maßnahmen', 'Herkunft', 'Evidenz', 'Freigabe', 'Hinweis']
BR_R = [7, 16, 12, 34, 9, 46, 30, 13, 46, 46, 22, 13, 15, 34]
ws = blatt(wb, 'Regeln', SP_R, BR_R)
for kat, seed in KATALOGE:
    for r in sorted(seed['rules'], key=lambda x: (x['hazard'], -x['priority'])):
        h = H[kat].get(r['hazard'], {})
        fuell = STUFE_FILL.get(r['result'])
        zeile(ws, [
            kat,
            r['code'],
            r['hazard'],
            h.get('title', ''),
            r['priority'],
            regelBedingung(kat, r),
            ausdruck(kat, r.get('applicability'), 0),
            STUFE.get(r['result'], r['result']),
            massnahmen(kat, r, 'sofort'),
            massnahmen(kat, r, 'mittel'),
            HERKUNFT.get(r.get('origin'), r.get('origin', '')),
            EVIDENZ.get(r.get('evidence'), r.get('evidence', '')),
            FREIGABE.get(r.get('quality_status'), r.get('quality_status', '')),
            (r.get('notes') or '').strip(),
        ], BR_R, PatternFill('solid', fgColor=fuell) if fuell else None)
abschluss(ws, len(SP_R))

# --- Karten -----------------------------------------------------------------
SP_K = ['Katalog', 'Nr.', 'Karte', 'Erhebungsbereich', 'Art', 'Positionen', 'Hilfetext']
BR_K = [7, 8, 34, 26, 12, 76, 44]
ART = {'checklist': 'Ankreuzliste', 'card': 'gemischte Karte'}
ws = blatt(wb, 'Karten', SP_K, BR_K)
for kat, seed in KATALOGE:
    for g in seed.get('question_groups', []):
        positionen = []
        for it in g.get('items', []):
            nr = frageName(kat, it['question'], True)
            ort = ('%s: ' % it['row']) if it.get('row') else ''
            if it.get('mode') == 'check':
                # Ankreuzposition: Häkchen setzt den auffälligen Wert, das
                # Bestätigen der Karte den unauffälligen.
                txt = '%s„%s" → %s = %s' % (ort, it.get('label', ''), nr,
                                            wertText(kat, it['question'], it.get('value')))
                if it.get('clear') is not None:
                    txt += ' (bestätigt: %s)' % wertText(kat, it['question'], it['clear'])
                if it.get('implies'):
                    txt += ' [setzt mit: %s]' % ', '.join(
                        '%s = %s' % (frageName(kat, im['question'], True),
                                     wertText(kat, im['question'], im['value']))
                        for im in it['implies'])
            else:
                # Native Position: die Frage erscheint unverändert auf der Karte.
                txt = '%s%s (Frage unverändert)' % (ort, frageName(kat, it['question']))
            positionen.append('• ' + txt)
        zeile(ws, [kat, g.get('ui_number', ''), g['title'], g.get('category', ''),
                   ART.get(g.get('kind'), g.get('kind', '')),
                   '\n'.join(positionen), g.get('help', '')], BR_K)
abschluss(ws, len(SP_K))

# --- Annahmen ---------------------------------------------------------------
SP_A = ['Frage', 'Angenommener Wert', 'Wenn', 'Frühestes Baujahr', 'Begründung']
BR_A = [52, 18, 40, 16, 86]
ws = blatt(wb, 'Annahmen', SP_A, BR_A)
for a in MF.get('assumptions', []):
    zeile(ws, [
        frageName('GBU', a['question']),
        wertText('GBU', a['question'], a.get('value')),
        ausdruck('GBU', a.get('when'), 0),
        a.get('min_baujahr') or '',
        a.get('reason', ''),
    ], BR_A)
if MF.get('assumptions_void_when'):
    zeile(ws, ['Alle Annahmen entfallen, wenn:', '',
               ausdruck('GBU', MF['assumptions_void_when'], 0), '',
               'Ohne geprüften Konformitätsnachweis trägt das Baujahr keine '
               'Vermutung mehr – dann ist zu erheben.'], BR_A, ZWISCHEN_FILL)
abschluss(ws, len(SP_A))

# --- Nachweise --------------------------------------------------------------
SP_N = ['Frage', 'Vorbelegter Wert', 'Wenn', 'Frühestes Baujahr', 'Begründung',
        'Quelle']
BR_N = [52, 18, 40, 16, 76, 30]
ws = blatt(wb, 'Nachweise', SP_N, BR_N)
for n in MF.get('nachweise', []):
    zeile(ws, [
        frageName('GBU', n['question']),
        wertText('GBU', n['question'], n.get('value')),
        ausdruck('GBU', n.get('when'), 0),
        n.get('min_baujahr') or '',
        n.get('reason', ''),
        n.get('source', ''),
    ], BR_N)
abschluss(ws, len(SP_N))

# --- Maßnahmen --------------------------------------------------------------
SP_M = ['Katalog', 'Code', 'Maßnahme', 'Art', 'Gruppe', 'In Regeln']
BR_M = [7, 18, 86, 16, 12, 10]
ws = blatt(wb, 'Maßnahmen', SP_M, BR_M)
for kat, seed in KATALOGE:
    verwendet = {}
    for r in seed['rules']:
        for mm in r.get('measures', []):
            verwendet.setdefault(mm['measure'], []).append(r['code'])
    for m in seed['measures']:
        zeile(ws, [
            kat, m['code'], m.get('title', ''),
            TYP.get(m.get('type'), m.get('type', '')),
            'Sofort' if m.get('priority_class') == 'SOFORT' else 'mittelfristig',
            len(verwendet.get(m['code'], [])),
        ], BR_M)
abschluss(ws, len(SP_M))

wb.save(OUT)
print('geschrieben: %s' % OUT)
print('  Fragen      GBU %3d / Cyber %3d' % (len(MF['questions']), len(CY['questions'])))
print('  Gefährdungen GBU %3d / Cyber %3d' % (len(MF['hazards']), len(CY['hazards'])))
print('  Regeln      GBU %3d / Cyber %3d' % (len(MF['rules']), len(CY['rules'])))
print('  Karten      GBU %3d / Cyber %3d' % (len(MF.get('question_groups', [])),
                                            len(CY.get('question_groups', []))))
print('  Annahmen %d, Nachweise %d'
      % (len(MF.get('assumptions', [])), len(MF.get('nachweise', []))))
