# -*- coding: utf-8 -*-
"""Regelprüfung des MF-Katalogs (EN 81-20, mehrfragig): legt alle Regeln mit
quality_status = REVIEW_REQUIRED zur fachlichen Freigabe vor.

    python3 gen_mf_regelpruefung_xlsx.py            # alle offenen Regeln
    python3 gen_mf_regelpruefung_xlsx.py --out X.xlsx

Eingabe: norm_81_20_mf.json (+ norm_81_80_mf.json, en8180_map.json,
         pruefbericht_*.json, falls vorhanden)
Ausgabe: GBU_MF_Regelpruefung.xlsx
Rückweg: python3 apply_mf_regelpruefung.py  (schreibt mf_content/regelfreigabe.py
         und mf_content/massnahmenart.py)

Aufbau (Stand 16.09.2026, dritte Fassung nach zwei Prüfrunden):
  * Lesehinweise
  * Übersicht – Fortschritt je Erhebungsbereich (zählt jede Regel mit
    wirksamer Entscheidung, Korrektur-fehlt-Zähler)
  * Prüfbericht – Bewertung jedes Befunds des externen Prüfberichts und was
    daraus geworden ist; offene Punkte mit Entscheidungsspalte
  * ein Blatt je Erhebungsbereich – je Gefährdung ein Block:
      Kopfzeile  (Art „Gefährdung") mit Anwendbarkeit, Pflichtfragen,
                  EN-81-80-Bezug und Sammelentscheidung,
      Regelzeilen (Art „Regel") – eine Einzelentscheidung geht vor.
    Freigegebene Regeln stehen grau im Block; nach dem Prüfbericht
    geänderte Regeln sind orange markiert.
  * Maßnahmenarten – T/O je Maßnahme zum Gegenlesen
  * Fragen – alle Fragen der gezeigten Gefährdungen mit Sichtbarkeit und
    Pflicht

Die Spaltenlage ist für apply_mf_regelpruefung.py verbindlich (SPALTEN,
MA_SPALTEN) – bei Änderungen beide Dateien anpassen.
"""
import glob, json, os, sys
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
argv = sys.argv[1:]
OUT = os.path.join(HERE, 'GBU_MF_Regelpruefung.xlsx')
if '--out' in argv:
    OUT = argv[argv.index('--out') + 1]
STAND = '16.09.2026'


def lade(name):
    p = os.path.join(HERE, name)
    return json.load(open(p, encoding='utf-8')) if os.path.exists(p) else None


seed = lade('norm_81_20_mf.json')
en8180 = lade('norm_81_80_mf.json')
IN_8180 = {h['code'] for h in en8180['hazards']} if en8180 else None
EN_PUNKTE = {}
for p in (lade('en8180_map.json') or {}).get('punkte', []):
    for hc in p.get('hazards', []):
        EN_PUNKTE.setdefault(hc, []).append('Nr. %s (%s)' % (p['nr'], p['abschnitt']))
def _ist_regelbericht(d):
    """Blatt „Prüfbericht" erwartet Befunde als Tupel (id, prio, bereich, …).

    Seit 20.09.2026 liegt neben pruefbericht_<Datum>.json auch
    pruefbericht_riedl_<Datum>.json im Ordner – der Prüfbericht zum
    FRAGENKATALOG, mit benannten Feldern und einem eigenen Generator
    (gen_riedl_pruefbericht_klaerung_xlsx.py). Der alphabetisch letzte Treffer
    war damit der falsche und der Generator brach ab. Jetzt entscheidet die
    Form, nicht der Dateiname.
    """
    if not isinstance(d, dict):
        return False
    runden = d.get('runden') or ([{'befunde': d['befunde']}] if 'befunde' in d else [])
    for rd in runden:
        for b in rd.get('befunde', []):
            return isinstance(b, (list, tuple))
    return False


BERICHT = None
for _pfad in sorted(glob.glob(os.path.join(HERE, 'pruefbericht_*.json'))):
    _d = json.load(open(_pfad, encoding='utf-8'))
    if _ist_regelbericht(_d):
        BERICHT = _d

Q = {q['code']: q for q in seed['questions']}
H = {h['code']: h for h in seed['hazards']}
IDS = (lade(os.path.join('mf_content', 'regel_ids.json')) or {}).get('regeln', {})


def revision(r):
    return IDS.get(r['hazard'], {}).get(r['code'], {}).get('rev', 1)
M = {m['code']: m for m in seed['measures']}
PB = '[Prüfbericht'

SPALTEN = ['Art', 'Code', 'Prio', 'Gefährdung / Bedingung', 'Stufe',
           'Sofortmaßnahme', 'Mittelfristige Maßnahme', 'Hinweis / Quelle',
           'Status', 'Entscheidung', 'Korrektur', 'wirksam']
BREITEN = [11, 12, 6, 72, 12, 38, 38, 38, 18, 14, 40, 12]
COL_ART, COL_CODE, COL_STATUS, COL_ENT, COL_KORR, COL_WIRK = 1, 2, 9, 10, 11, 12
MA_SPALTEN = ['Code', 'Maßnahme', 'Zeitlage', 'Art (neu)', 'Art bisher', 'geändert',
              'verwendet in', 'Art richtig?']

LAB = {'HIGH': 'Hoch', 'MEDIUM': 'Mittel', 'LOW': 'Niedrig',
       'NO_RISK': 'Kein Risiko', 'NOT_APPLICABLE': 'Nicht anwendbar'}
TYP = {'TECHNICAL': 'T', 'ORGANISATIONAL': 'O', 'PERSONAL': 'P'}
LOGIK = {
    'NONE': 'Vorrang – es gilt die zutreffende Regel mit der HÖCHSTEN PRIORITÄT '
            '(so wirken Kompensationen); bei gleicher Priorität die höhere Stufe.',
    'MAXIMUM': 'Höchste Stufe – alle zutreffenden Regeln zählen, maßgeblich ist die '
               'HÖCHSTE STUFE; Maßnahmen aller Treffer werden übernommen.',
    'ANY': 'Ortsmatrix – jede zutreffende Regel ist ein Befund, maßgeblich ist die '
           'HÖCHSTE STUFE.',
}

FONT = 'Arial'
def f(bold=False, color=None, size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
HEAD_FILL = PatternFill('solid', fgColor='1F3A5F')
HAZ_FILL = PatternFill('solid', fgColor='DCE6F1')
INPUT_FILL = PatternFill('solid', fgColor='FFF2CC')
DONE_FILL = PatternFill('solid', fgColor='EDEDED')
PB_FILL = PatternFill('solid', fgColor='FAD7B5')
ROT_FILL = PatternFill('solid', fgColor='F4B6B6')
STUFE_FILL = {'HIGH': 'F4CCCC', 'MEDIUM': 'FCE5CD', 'LOW': 'FFF2CC', 'NO_RISK': 'D9EAD3'}
thin = Side(style='thin', color='C9D1DB')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
OFFEN = Protection(locked=False)


# ---------------------------------------------------------------- Texte
def val_label(qc, v):
    q = Q.get(qc)
    if q and q.get('type') == 'SELECT':
        for o in q.get('options', []):
            if o['value'] == v:
                return '„%s"' % o['label']
    if v is True: return 'Ja'
    if v is False: return 'Nein'
    return str(v).replace('.', ',') if isinstance(v, float) else str(v)


def qname(qc, kurz=False):
    q = Q.get(qc)
    if not q: return qc
    return q.get('ui_number', qc) if kurz else '%s %s' % (q.get('ui_number', ''), q['text'])


def expr(e, tiefe=0, kurz=False):
    if not e: return ''
    if 'all' in e:
        teile = [expr(x, tiefe + 1, kurz) for x in e['all']]
        if tiefe == 0 and not kurz:
            return '\nUND '.join(teile)
        return '(' + ' UND '.join(teile) + ')'
    if 'any' in e:
        teile = [expr(x, tiefe + 1, kurz) for x in e['any']]
        s = ' ODER '.join(teile)
        return s if tiefe == 0 else '(' + s + ')'
    if 'not' in e: return 'NICHT ' + expr(e['not'], tiefe + 1, kurz)
    qc, op, v = e.get('question'), e.get('operator'), e.get('value')
    n = qname(qc, kurz)
    if op == 'ANSWERED': return '[%s] ist beantwortet' % n
    if op == 'NOT_ANSWERED': return '[%s] ist unbeantwortet' % n
    ops = {'EQ': '=', 'NEQ': '≠', 'GT': '>', 'GTE': '≥', 'LT': '<',
           'LTE': '≤', 'IN': 'ist eines von', 'NOT_IN': 'ist keines von'}
    if isinstance(v, list):
        return '[%s] %s %s' % (n, ops.get(op, op), ' / '.join(val_label(qc, x) for x in v))
    return '[%s] %s %s' % (n, ops.get(op, op), val_label(qc, v))


def measures_of(r):
    sofort, mittel = [], []
    for mm in r.get('measures', []):
        m = M.get(mm['measure'])
        if not m: continue
        text = '• %s [%s]' % (m.get('title', ''), TYP.get(m.get('type'), '?'))
        (sofort if mm.get('group_id') == 'sofort' else mittel).append(text)
    return '\n'.join(sofort), '\n'.join(mittel)


def sources_of(h):
    return '\n'.join(('%s %s' % (s.get('document', ''), s.get('section', ''))).strip()
                     for s in h.get('sources', []))


def anwendbarkeit(h):
    teile = []
    for x in h.get('questions', []):
        if x['role'] != 'APPLICABILITY':
            continue
        if x.get('applicable_when'):
            teile.append(expr(x['applicable_when'], 1, kurz=False))
        elif not any(y.get('applicable_when') for y in h['questions'] if y['role'] == 'APPLICABILITY'):
            teile.append('[%s] = Ja' % qname(x['question']))
    return ' UND '.join(dict.fromkeys(teile)) if teile else 'immer'


def pflichtfragen(h):
    teile = []
    for x in h.get('questions', []):
        mode = x.get('required_mode', 'NEVER')
        if mode == 'ALWAYS':
            teile.append(qname(x['question'], True))
        elif mode == 'CONDITIONAL':
            teile.append('%s (wenn %s)' % (qname(x['question'], True),
                                           expr(x.get('required_when'), 1, kurz=True)))
    return ', '.join(teile) or '–'


def ist_auffang(r):
    return r['result'] == 'NO_RISK' and 'keine Mangelregel' in (r.get('notes') or '')


def hinweis(r):
    teile = []
    if ist_auffang(r):
        teile.append('Auffangregel (Mechanik): greift, wenn keine Mangelregel zutrifft '
                     'und alle Pflichtfragen beantwortet sind.')
    else:
        n = (r.get('notes') or '').strip()
        if n: teile.append(n)
    if r.get('evidence') == 'INFERRED':
        teile.append('Evidenz: abgeleitet')
    if r.get('applicability'):
        teile.append('Nur anwendbar, wenn: ' + expr(r['applicability']))
    gruppen = {m.get('group_id') for m in r.get('measures', [])}
    if r['result'] in ('HIGH', 'MEDIUM', 'LOW'):
        if 'sofort' not in gruppen:
            teile.append('AUFFÄLLIG: ohne Sofortmaßnahme')
        if not (gruppen - {'sofort'}):
            teile.append('AUFFÄLLIG: ohne mittelfristige Maßnahme')
    return '\n'.join(teile)


def zeilenhoehe(werte, breiten):
    zeilen = 1
    for wert, breite in zip(werte, breiten):
        if wert is None or (isinstance(wert, str) and wert.startswith('=')): continue
        s = str(wert)
        n = sum(max(1, -(-len(teil) // max(1, int(breite * 1.15)))) for teil in s.split('\n'))
        zeilen = max(zeilen, n)
    return min(409, 13.5 * zeilen + 3)


def kopf(ws, spalten, breiten, hidden=()):
    for i, (c, w) in enumerate(zip(spalten, breiten), 1):
        z = ws.cell(row=1, column=i, value=c)
        z.font = f(bold=True, color='FFFFFF'); z.fill = HEAD_FILL
        z.alignment = WRAP; z.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
        if i in hidden:
            ws.column_dimensions[get_column_letter(i)].hidden = True
    ws.freeze_panes = 'A2'


def schuetzen(ws):
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def liste(werte, titel):
    dv = DataValidation(type='list', formula1='"%s"' % ','.join(werte), allow_blank=True,
                        showErrorMessage=True, errorStyle='stop', errorTitle='Ungültige Eingabe',
                        error='Bitte einen Wert aus der Liste wählen: ' + ', '.join(werte),
                        promptTitle=titel[:32], prompt='Auswahl: ' + ', '.join(werte),
                        showInputMessage=True)
    return dv


# ---------------------------------------------------------------- Daten
offen = [r for r in seed['rules'] if r.get('quality_status') == 'REVIEW_REQUIRED']
offen_h = {r['hazard'] for r in offen}
rules_by_h = {}
for r in seed['rules']:
    rules_by_h.setdefault(r['hazard'], []).append(r)

bereiche = []
for h in seed['hazards']:
    g = h.get('ui_group', '?')
    if h['code'] in offen_h and g not in bereiche:
        bereiche.append(g)

BLATT = {'A': 'A Anlagenmerkmale', 'Z': 'Z Zugang', 'M': 'M Triebwerksraum',
         'T': 'T Türen', 'K': 'K Fahrkorb', 'F': 'F Fahrkorbdach',
         'S': 'S Schacht', 'G': 'G Schachtgrube', 'U': 'U Umfeld',
         'SF': 'SF Sonderfunktionen', 'D': 'D Unterlagen'}


def blattname(g):
    return BLATT.get(g.split(' – ')[0].strip(), g)[:31]


wb = Workbook()

# ---------------------------------------------------------------- Lesehinweise
ws = wb.active
ws.title = 'Lesehinweise'
ws.column_dimensions['A'].width = 120
n_auffang = sum(1 for r in offen if ist_auffang(r))
n_pb = sum(1 for r in offen if PB in (r.get('notes') or '') or 'früheren Stand' in (r.get('notes') or ''))
texte = [
    ('Regelprüfung GBU EN 81-20 (Fragebogen) – offene Regeln', True),
    ('Regelwerk %s · Stand %s · %d offene Regeln in %d Gefährdungen, davon %d Auffangregeln '
     '(Kein Risiko), %d fachliche Regeln und %d nach dem externen Prüfbericht geänderte Regeln.'
     % (seed['rule_version'], STAND, len(offen), len(offen_h), n_auffang,
        len(offen) - n_auffang, n_pb), False),
    ('', False),
    ('Zweite Fassung: Der externe Prüfbericht zur ersten Fassung ist eingearbeitet. Blatt '
     '„Prüfbericht“ zeigt je Befund, ob er zutrifft und was geändert wurde; die geänderten Regeln '
     'sind in den Bereichsblättern orange markiert (Status „offen – geändert“).', False),
    ('Der Katalog „Bestand nach DIN EN 81-80“ ist aus diesem Katalog abgeleitet – eine Freigabe hier '
     'gilt dort automatisch. Das ist zulässig, weil der Generator Regeln, Gefährdungen und '
     'Anwendbarkeit dort unverändert übernimmt und bei jeder Abweichung abbricht. Die Kopfzeile nennt '
     'die zugehörigen EN-81-80-Punkte bzw. „nicht im Bestandskatalog“.', False),
    ('Dritte Fassung (16.09.2026): Regel-IDs sind jetzt dauerhaft – eine ID bezeichnet immer denselben '
     'Sachverhalt; neue Sachverhalte bekommen neue Nummern (deshalb Lücken). Die Prüfliste v2 ist '
     'damit überholt, bitte nur diese Datei ausfüllen. „Rev.“ im Status zählt inhaltliche Änderungen '
     'seit der ersten Fassung.', False),
    ('Baujahr-Annahmen (Festlegung 07.09., bestätigt 16.09.2026): Ab dem genannten Baujahr gilt ein '
     'Ausstattungsmerkmal als vorhanden, solange niemand widerspricht; im Bericht steht dann '
     '„angenommen“. Blatt „Fragen“, Spalte „Annahme ab Baujahr“ zeigt, welche Fragen betroffen sind.', False),
    ('', False),
    ('So wird entschieden', True),
    ('1. Je Erhebungsbereich ein Blatt, darin je Gefährdung ein Block. Die blaue Kopfzeile nennt '
     'Titel, Faktor, betroffene Personen, Bewertungslogik, wann die Gefährdung überhaupt bewertet '
     'wird (Anwendbarkeit), welche Fragen vorher beantwortet sein müssen (Pflichtfragen), '
     'EN-81-80-Bezug und Quellen.', False),
    ('2. Passt der ganze Block: in der Kopfzeile „Freigeben“ wählen – das gilt für alle offenen '
     'Regeln der Gefährdung, einschließlich der Auffangregel.', False),
    ('3. Passt eine einzelne Regel nicht: in ihrer Zeile „Ändern“ oder „Streichen“ wählen und in '
     '„Korrektur“ sagen, was gelten soll. Ohne Korrektur wird die Zelle rot. Die '
     'Einzelentscheidung geht der Sammelentscheidung vor.', False),
    ('4. Leer lassen = noch nicht entschieden. Blockweise arbeiten und zwischendurch zurückgeben '
     'ist möglich.', False),
    ('5. Blatt „Prüfbericht“: Bei den offenen Punkten (Spalte „Entscheidung nötig“) bitte '
     'Übernehmen / Ablehnen / Später wählen und ggf. kommentieren.', False),
    ('6. Blatt „Maßnahmenarten“: Die Einstufung technisch/organisatorisch wurde neu abgeleitet. '
     'Falsche Einstufungen in „Art richtig?“ mit T oder O korrigieren; leer = passt.', False),
    ('', False),
    ('Worauf achten', True),
    ('• Stufe: Trifft Hoch / Mittel / Niedrig für diesen Sachverhalt zu?', False),
    ('• Bedingung: Ist der Sachverhalt richtig getroffen – auch UND/ODER und die Ausnahmen (NICHT)?', False),
    ('• Priorität / Bewertungslogik: Bei „Vorrang“ gewinnt die Regel mit der höchsten Prio – dort '
     'stecken die Kompensationen. Bei „Höchste Stufe“ zählen alle zutreffenden Regeln.', False),
    ('• Maßnahmen: sofort und mittelfristig; [T] technisch, [O] organisatorisch. „AUFFÄLLIG“ im '
     'Hinweis markiert fehlende Maßnahmen.', False),
    ('• Graue Zeilen sind bereits freigegeben und stehen nur zum Verständnis im Block.', False),
    ('• Auffangregeln (Kein Risiko, Prio 1) sind Mechanik: Sie greifen erst, wenn keine Mangelregel '
     'zutrifft und alle Pflichtfragen beantwortet sind; vorher meldet die Engine „unvollständig“.', False),
    ('', False),
    ('Die Blätter sind ohne Kennwort geschützt, damit nichts versehentlich überschrieben wird; '
     'offen sind nur die gelben Eingabezellen. Filter funktionieren. Aufheben: Überprüfen → '
     'Blattschutz aufheben.', False),
    ('', False),
    ('Rückweg (technisch): python3 apply_mf_regelpruefung.py → python3 gen_mf_catalog.py → '
     'python3 gen_en8180_catalog.py → App-Assets erneuern → Engine neu seeden. Freigaben sind an den '
     'Regelinhalt gebunden: Wird eine Regel später geändert, fällt sie automatisch wieder auf „offen“.', False),
]
for i, (t, b) in enumerate(texte, 1):
    z = ws.cell(row=i, column=1, value=t)
    z.font = f(bold=b, size=13 if i == 1 else (11 if b else 10))
    z.alignment = WRAP
    if not b and t:
        ws.row_dimensions[i].height = 14 * max(1, -(-len(t) // 125)) + 2
schuetzen(ws)

# ---------------------------------------------------------------- Übersicht
wsu = wb.create_sheet('Übersicht')
UKOPF = ['Erhebungsbereich', 'Blatt', 'Gefährdungen', 'offene Regeln', 'davon Auffang',
         'davon geändert (Prüfbericht)', 'entschieden', 'Freigeben', 'Ändern', 'Streichen',
         'Korrektur fehlt', 'Fortschritt', 'fachlich freigabereif']
kopf(wsu, UKOPF, [44, 20, 13, 13, 12, 16, 13, 12, 10, 10, 12, 12, 14])

# ---------------------------------------------------------------- Prüfbericht
if BERICHT:
    wsp = wb.create_sheet('Prüfbericht')
    PKOPF = ['ID', 'Priorität', 'Bereich / Regeln', 'Bewertung', 'Begründung / Beleg',
             'Umsetzung', 'Entscheidung nötig', 'Deine Entscheidung', 'Kommentar']
    kopf(wsp, PKOPF, [7, 10, 22, 16, 60, 60, 12, 16, 40])
    dvp = liste(['Übernehmen', 'Ablehnen', 'Später'], 'Entscheidung')
    wsp.add_data_validation(dvp)
    BEW_FILL = {'bestätigt': 'D9EAD3', 'teilweise': 'FFF2CC'}
    runden = BERICHT.get('runden') or [{'titel': BERICHT.get('titel', ''), 'befunde': BERICHT['befunde']}]
    alle = []
    for rd in runden:
        alle.append(None)
        alle.append(('__titel__', rd['titel']))
        alle.extend(rd['befunde'])
    ri = 1
    for b in alle:
        ri += 1
        if b is None:
            ri -= 1
            continue
        if b[0] == '__titel__':
            z = wsp.cell(row=ri, column=1, value=b[1])
            z.font = f(bold=True, size=11); z.fill = HAZ_FILL
            wsp.merge_cells(start_row=ri, start_column=1, end_row=ri, end_column=9)
            continue
        bid, prio, ber, bew, grund, um, noetig = b
        werte = [bid, prio, ber, bew, grund, um, 'ja' if noetig else '', None, None]
        for ci, wert in enumerate(werte, 1):
            z = wsp.cell(row=ri, column=ci, value=wert)
            z.font = f(bold=ci == 1); z.alignment = WRAP; z.border = BORDER
            if ci == 4:
                z.fill = PatternFill('solid', fgColor=BEW_FILL.get(bew.split(' ')[0], 'EDEDED'))
            if ci in (8, 9) and noetig:
                z.fill = INPUT_FILL; z.protection = OFFEN
        if noetig:
            dvp.add(wsp.cell(row=ri, column=8))
        wsp.row_dimensions[ri].height = zeilenhoehe(werte, [7, 10, 22, 16, 60, 60, 12, 16, 40])
    schuetzen(wsp)

# ---------------------------------------------------------------- Bereichsblätter
fragen = []


def sammle(e):
    if not e: return
    if 'all' in e:
        for x in e['all']: sammle(x)
    elif 'any' in e:
        for x in e['any']: sammle(x)
    elif 'not' in e:
        sammle(e['not'])
    elif e.get('question') and e['question'] not in fragen:
        fragen.append(e['question'])
        sammle(Q.get(e['question'], {}).get('visible_when'))


verwendet = {}   # Maßnahme -> Regeln (nur gezeigte Gefährdungen)
E, A, S, K, W = (get_column_letter(c) for c in (COL_ENT, COL_ART, COL_STATUS, COL_KORR, COL_WIRK))

for bi, g in enumerate(bereiche):
    name = blattname(g)
    ws = wb.create_sheet(name)
    kopf(ws, SPALTEN, BREITEN, hidden=(COL_WIRK,))
    dv_h = liste(['Freigeben'], 'Sammelentscheidung')
    dv_r = liste(['Freigeben', 'Ändern', 'Streichen'], 'Entscheidung')
    ws.add_data_validation(dv_h); ws.add_data_validation(dv_r)

    zeile = 2
    n_h = n_r = n_a = n_g = 0
    for h in seed['hazards']:
        if h.get('ui_group') != g or h['code'] not in offen_h:
            continue
        n_h += 1
        regeln = sorted(rules_by_h[h['code']], key=lambda r: -r['priority'])
        n_off = sum(1 for r in regeln if r.get('quality_status') == 'REVIEW_REQUIRED')
        agg = h.get('aggregation_type', 'NONE')
        for x in h.get('questions', []):
            sammle({'question': x['question']})
            sammle(x.get('required_when')); sammle(x.get('applicable_when'))
        kopf_text = ('%s\nFaktor: %s\nBetroffen: %s\nBewertungslogik: %s\n'
                     'Bewertet, wenn: %s\nPflichtfragen: %s') % (
            h['title'], h.get('hazard_factor', '–'),
            ', '.join(h.get('person_groups', [])) or '–', LOGIK.get(agg, agg),
            anwendbarkeit(h), pflichtfragen(h))
        if h.get('description'):
            kopf_text += '\n' + h['description']
        if IN_8180 is None:
            auch = ''
        elif h['code'] in IN_8180:
            auch = 'EN 81-80: ' + (', '.join(EN_PUNKTE.get(h['code'], [])) or 'im Bestandskatalog (organisatorisch)')
        else:
            auch = 'EN 81-80: nicht im Bestandskatalog (nur EN 81-20)'
        hz_zeile = zeile
        werte = ['Gefährdung', h['code'], None, kopf_text, None, None, None,
                 (auch + '\n' + sources_of(h)).strip(), '%d von %d offen' % (n_off, len(regeln)),
                 None, None, None]
        for ci, wert in enumerate(werte, 1):
            z = ws.cell(row=zeile, column=ci, value=wert)
            z.font = f(bold=ci in (1, 2))
            z.fill = INPUT_FILL if ci == COL_ENT else HAZ_FILL
            z.alignment = WRAP; z.border = BORDER
        ws.cell(row=zeile, column=COL_ENT).protection = OFFEN
        ws.row_dimensions[zeile].height = zeilenhoehe(werte, BREITEN)
        dv_h.add(ws.cell(row=zeile, column=COL_ENT))
        zeile += 1

        for r in regeln:
            ist_offen = r.get('quality_status') == 'REVIEW_REQUIRED'
            geaendert = PB in (r.get('notes') or '') or 'früheren Stand' in (r.get('notes') or '')
            sofort, mittel = measures_of(r)
            for mm in r.get('measures', []):
                verwendet.setdefault(mm['measure'], []).append(r['code'])
            if not ist_offen:
                status = 'freigegeben'
            elif ist_auffang(r):
                status = 'offen (Auffang)'; n_a += 1
            elif geaendert:
                status = 'offen – geändert (Rev. %d)' % revision(r); n_g += 1
            else:
                status = 'offen'
            if ist_offen:
                n_r += 1
            sammle(r['condition']); sammle(r.get('applicability'))
            wirk = ('=IF(%s%d<>"",%s%d,IF($%s$%d="Freigeben","Freigeben",""))'
                    % (E, zeile, E, zeile, E, hz_zeile)) if ist_offen else None
            werte = ['Regel', r['code'], r['priority'], expr(r['condition']),
                     LAB.get(r['result'], r['result']), sofort, mittel, hinweis(r),
                     status, None, None, wirk]
            for ci, wert in enumerate(werte, 1):
                z = ws.cell(row=zeile, column=ci, value=wert)
                z.font = f(color=None if ist_offen else '808080')
                z.alignment = WRAP; z.border = BORDER
                if not ist_offen:
                    z.fill = DONE_FILL
                elif ci in (COL_ENT, COL_KORR):
                    z.fill = INPUT_FILL; z.protection = OFFEN
                elif ci == 5:
                    z.fill = PatternFill('solid', fgColor=STUFE_FILL.get(r['result'], 'FFFFFF'))
                elif geaendert and ci in (COL_CODE, COL_STATUS):
                    z.fill = PB_FILL
            ws.row_dimensions[zeile].height = zeilenhoehe(werte, BREITEN)
            if ist_offen:
                dv_r.add(ws.cell(row=zeile, column=COL_ENT))
            zeile += 1
    letzte = zeile - 1
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(SPALTEN)), letzte)
    ws.conditional_formatting.add(
        '%s2:%s%d' % (K, K, letzte),
        FormulaRule(formula=['AND(OR(%s2="Ändern",%s2="Streichen"),%s2="")' % (E, E, K)],
                    fill=ROT_FILL, stopIfTrue=True))
    schuetzen(ws)

    # Übersichtszeile mit Live-Zählung
    ur = bi + 2
    rng = lambda col: "'%s'!$%s$2:$%s$%d" % (name, col, col, letzte)
    werte = [g, name, n_h, n_r, n_a, n_g,
             '=COUNTIF(%s,"?*")' % rng(W),
             '=COUNTIF(%s,"Freigeben")' % rng(W),
             '=COUNTIF(%s,"Ändern")' % rng(W),
             '=COUNTIF(%s,"Streichen")' % rng(W),
             '=COUNTIFS(%s,"Ändern",%s,"")+COUNTIFS(%s,"Streichen",%s,"")'
             % (rng(E), rng(K), rng(E), rng(K)),
             '=IF(D%d=0,"",G%d/D%d)' % (ur, ur, ur),
             '=IF(AND(H%d=D%d,K%d=0),"ja","nein")' % (ur, ur, ur)]
    for ci, wert in enumerate(werte, 1):
        z = wsu.cell(row=ur, column=ci, value=wert)
        z.font = f(); z.border = BORDER; z.alignment = WRAP
    wsu.cell(row=ur, column=12).number_format = '0%'

last = len(bereiche) + 1
sumrow = last + 1
wsu.cell(row=sumrow, column=1, value='Summe').font = f(bold=True)
for ci in range(3, 12):
    col = get_column_letter(ci)
    z = wsu.cell(row=sumrow, column=ci, value='=SUM(%s2:%s%d)' % (col, col, last))
    z.font = f(bold=True); z.border = BORDER
z = wsu.cell(row=sumrow, column=12, value='=IF(D%d=0,"",G%d/D%d)' % (sumrow, sumrow, sumrow))
z.number_format = '0%'; z.font = f(bold=True); z.border = BORDER
z = wsu.cell(row=sumrow, column=13, value='=IF(AND(H%d=D%d,K%d=0),"ja","nein")' % (sumrow, sumrow, sumrow))
z.font = f(bold=True); z.border = BORDER
wsu.conditional_formatting.add('M2:M%d' % sumrow,
                               FormulaRule(formula=['M2="ja"'], fill=PatternFill('solid', fgColor='D9EAD3')))
wsu.conditional_formatting.add('K2:K%d' % sumrow,
                               FormulaRule(formula=['K2>0'], fill=ROT_FILL))
wsu.cell(row=sumrow + 2, column=1,
         value='„entschieden“ zählt jede offene Regel mit wirksamer Entscheidung – der eigenen oder '
               'der Sammelfreigabe ihrer Gefährdung. „fachlich freigabereif“ = jede offene Regel ist '
               'freigegeben und keine Korrektur fehlt; „Ändern“/„Streichen“ bleiben offen, bis der '
               'Katalog angepasst und erneut vorgelegt ist.'
         ).font = f(italic=True, color='555555')
schuetzen(wsu)

# ---------------------------------------------------------------- Maßnahmenarten
wsm = wb.create_sheet('Maßnahmenarten')
kopf(wsm, MA_SPALTEN, [14, 70, 12, 10, 10, 10, 40, 12])
dvm = liste(['T', 'O'], 'Art richtig?')
wsm.add_data_validation(dvm)
ri = 2
for code, regeln in sorted(verwendet.items(), key=lambda kv: (M[kv[0]]['priority_class'], M[kv[0]]['title'])):
    m = M[code]
    bisher = 'O' if m['priority_class'] == 'SOFORT' else 'T'
    neu = TYP.get(m['type'], '?')
    werte = [code, m['title'], 'sofort' if m['priority_class'] == 'SOFORT' else 'mittelfristig',
             neu, bisher, 'ja' if neu != bisher else '', ', '.join(sorted(set(regeln))), None]
    for ci, wert in enumerate(werte, 1):
        z = wsm.cell(row=ri, column=ci, value=wert)
        z.font = f(bold=ci == 4); z.alignment = WRAP; z.border = BORDER
        if ci == 8:
            z.fill = INPUT_FILL; z.protection = OFFEN
        elif ci == 6 and neu != bisher:
            z.fill = PB_FILL
    dvm.add(wsm.cell(row=ri, column=8))
    wsm.row_dimensions[ri].height = zeilenhoehe(werte, [14, 70, 12, 10, 10, 10, 40, 12])
    ri += 1
wsm.auto_filter.ref = 'A1:H%d' % (ri - 1)
schuetzen(wsm)

# ---------------------------------------------------------------- Fragen
wsq = wb.create_sheet('Fragen')
cols = ['Nr.', 'Code', 'Frage', 'Erhebungsbereich', 'Typ', 'Antwortoptionen',
        'Sichtbar, wenn', 'Pflicht in', 'Annahme ab Baujahr', 'Hilfetext']
breiten = [8, 26, 50, 26, 9, 44, 36, 30, 20, 60]
kopf(wsq, cols, breiten)
if seed.get('assumptions'):
    for a in seed['assumptions']:
        sammle(a.get('when'))
    sammle(seed.get('assumptions_void_when'))
pflicht_in = {}
for h in seed['hazards']:
    if h['code'] not in offen_h:
        continue
    for x in h.get('questions', []):
        if x.get('required_mode', 'NEVER') != 'NEVER':
            pflicht_in.setdefault(x['question'], []).append(
                h['code'] + ('' if x['required_mode'] == 'ALWAYS' else ' (bedingt)'))
annahme = {a['question']: a for a in seed.get('assumptions', [])}
ordnung = {q['code']: i for i, q in enumerate(seed['questions'])}
for ri, qc in enumerate(sorted(fragen, key=lambda c: ordnung.get(c, 9999)), 2):
    q = Q.get(qc, {})
    if q.get('type') == 'SELECT':
        opts = '\n'.join('• %s%s' % (o['label'], ' (unauffällig)' if o['value'] == q.get('best_case') else '')
                         for o in q.get('options', []))
    elif q.get('type') == 'YES_NO':
        opts = 'Ja / Nein' + (' (unauffällig: %s)' % ('Ja' if q['best_case'] else 'Nein')
                              if 'best_case' in q else '')
    else:
        opts = 'Zahl %s … %s' % (q.get('min', ''), q.get('max', ''))
    an = annahme.get(qc)
    werte = [q.get('ui_number', ''), qc, q.get('text', ''), q.get('category', ''),
             q.get('type', ''), opts, expr(q.get('visible_when'), 1) or 'immer',
             ', '.join(pflicht_in.get(qc, [])) or '–',
             ('%s: %s' % (expr(an['when'], 1, kurz=True), val_label(qc, an['value']))) if an else '',
             q.get('help_text') or '–']
    for ci, wert in enumerate(werte, 1):
        z = wsq.cell(row=ri, column=ci, value=wert)
        z.font = f(); z.alignment = WRAP; z.border = BORDER
    wsq.row_dimensions[ri].height = zeilenhoehe(werte, breiten)
wsq.auto_filter.ref = 'A1:J%d' % (len(fragen) + 1)
schuetzen(wsq)

wb.save(OUT)


def neu_berechnen(pfad):
    """Formelergebnisse speichern (Prüfbericht H14): Die Datei einmal durch
    LibreOffice rechnen lassen, sonst zeigen Betrachter ohne Rechenkern leere
    Kennzahlen. Ohne LibreOffice bleibt die Datei unverändert (Excel rechnet
    beim Öffnen selbst)."""
    import shutil, subprocess, tempfile
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        print('Hinweis: LibreOffice nicht gefunden – Formelwerte werden erst beim Öffnen berechnet.')
        return
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, '--headless', '--calc', '--convert-to',
                        'xlsx:Calc MS Excel 2007 XML', '--outdir', tmp, pfad],
                       check=False, capture_output=True, timeout=300)
        ziel = os.path.join(tmp, os.path.basename(pfad))
        if os.path.exists(ziel):
            shutil.copyfile(ziel, pfad)
            print('Formelwerte berechnet und gespeichert.')
        else:
            print('Hinweis: Neuberechnung fehlgeschlagen – Datei bleibt unberechnet.')


if '--ohne-berechnung' not in argv:
    neu_berechnen(OUT)
print('geschrieben: %s | %d offene Regeln (%d Auffang, %d geändert) | %d Gefährdungen | '
      '%d Bereiche | %d Fragen | %d Maßnahmen'
      % (OUT, len(offen), n_auffang, n_pb, len(offen_h), len(bereiche), len(fragen), len(verwendet)))
