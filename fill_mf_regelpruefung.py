# -*- coding: utf-8 -*-
"""Entscheidungen aus einer Prüfung in GBU_MF_Regelpruefung.xlsx eintragen.

    python3 fill_mf_regelpruefung.py entscheidungen.json [XLSX]

entscheidungen.json: Liste aus {"code", "entscheidung", "korrektur"}.
Geschrieben werden nur die Zeilen mit Art = „Regel" und passendem Code
(Spaltenlage wie in gen_mf_regelpruefung_xlsx.py: Art 1, Code 2,
Entscheidung 10, Korrektur 11). Danach:

    python3 apply_mf_regelpruefung.py --datum JJJJ-MM-TT

Der Umweg über die Datei ist Absicht: Arne sieht anschließend im gewohnten
Blatt, was entschieden wurde, und kann einzelne Zeilen überschreiben.
"""
import json
import os
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
COL_ART, COL_CODE, COL_ENT, COL_KORR = 1, 2, 10, 11
VALID = {'Freigeben', 'Ändern', 'Streichen'}

if len(sys.argv) < 2:
    sys.exit(__doc__)
quelle = sys.argv[1]
xlsx = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, 'GBU_MF_Regelpruefung.xlsx')

eintraege = json.load(open(quelle, encoding='utf-8'))
E = {}
for e in eintraege:
    ent = e['entscheidung']
    if ent not in VALID:
        sys.exit('unbekannte Entscheidung %r bei %s' % (ent, e['code']))
    if e['code'] in E and E[e['code']] != (ent, e.get('korrektur', '')):
        sys.exit('widersprüchliche Einträge für %s' % e['code'])
    E[e['code']] = (ent, (e.get('korrektur') or '').strip())

wb = load_workbook(xlsx)
getroffen, zaehl = set(), {}
for ws in wb.worksheets:
    kopf = [c.value for c in ws[1]]
    if not kopf or kopf[0] != 'Art' or len(kopf) < COL_KORR:
        continue
    for row in ws.iter_rows(min_row=2):
        if row[COL_ART - 1].value != 'Regel':
            continue
        code = row[COL_CODE - 1].value
        if code not in E:
            continue
        ent, korr = E[code]
        row[COL_ENT - 1].value = ent
        row[COL_KORR - 1].value = korr or None
        getroffen.add(code)
        zaehl[ent] = zaehl.get(ent, 0) + 1

fehlend = sorted(set(E) - getroffen)
wb.save(xlsx)
print('eingetragen: %d von %d (%s)'
      % (len(getroffen), len(E), ', '.join('%s %d' % kv for kv in sorted(zaehl.items()))))
if fehlend:
    print('nicht im Blatt gefunden (%d): %s' % (len(fehlend), ', '.join(fehlend)))
