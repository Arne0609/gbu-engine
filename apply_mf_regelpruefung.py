# -*- coding: utf-8 -*-
"""Übernimmt die zurückgegebene Regelprüfung des MF-Typs (EN 81-20):

    python3 apply_mf_regelpruefung.py [GBU_MF_Regelpruefung.xlsx] [--datum JJJJ-MM-TT] [--probe]

Liest die Bereichsblätter (Aufbau siehe gen_mf_regelpruefung_xlsx.py):
  * Kopfzeile „Gefährdung" mit „Freigeben" -> gilt für alle offenen Regeln
    dieser Gefährdung,
  * Zeile „Regel" mit Freigeben/Ändern/Streichen -> geht der Sammel-
    entscheidung vor.
und ergänzt mf_content/regelfreigabe.py. Bestehende Einträge bleiben
erhalten; eine neue Entscheidung zu einer noch offenen Regel ersetzt eine
ältere („Ändern" -> später „Freigeben").

Jeder neue Eintrag trägt Datum und Fingerabdruck des Regelinhalts
(Bedingung, Stufe, Priorität, Anwendbarkeit, Maßnahmen). Ändert sich die
Regel später im Inhalt oder verschiebt sich ihr Code, greift die Freigabe
nicht mehr und die Regel steht wieder auf REVIEW_REQUIRED.

--probe: nur auswerten und berichten, nichts schreiben.
Danach: python3 gen_mf_catalog.py && python3 gen_en8180_catalog.py
"""
import json, os, re, sys, datetime
from collections import Counter
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from mf_content.fingerabdruck import fingerabdruck  # noqa: E402

argv = sys.argv[1:]
datum = datetime.date.today().isoformat()
if '--datum' in argv:
    i = argv.index('--datum')
    datum = argv[i + 1]
    del argv[i:i + 2]
probe = '--probe' in argv
args = [a for a in argv if not a.startswith('--')]
xlsx = args[0] if args else os.path.join(HERE, 'GBU_MF_Regelpruefung.xlsx')

COL_ART, COL_CODE, COL_ENT, COL_KORR = 1, 2, 10, 11   # wie im Generator
VALID = {'Freigeben', 'Ändern', 'Streichen'}

seed = json.load(open(os.path.join(HERE, 'norm_81_20_mf.json'), encoding='utf-8'))
R = {r['code']: r for r in seed['rules']}
offen = {c for c, r in R.items() if r.get('quality_status') == 'REVIEW_REQUIRED'}


def text(v):
    return v.strip() if isinstance(v, str) else ''


wb = load_workbook(xlsx, data_only=True)
neu, warnungen = {}, []
for ws in wb.worksheets:
    kopf = [c.value for c in ws[1]]
    if not kopf or kopf[0] != 'Art' or len(kopf) < COL_KORR:
        continue
    sammel = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[COL_ART - 1]:
            continue
        art, code = row[COL_ART - 1], text(row[COL_CODE - 1])
        ent, korr = text(row[COL_ENT - 1]), text(row[COL_KORR - 1])
        if ent and ent not in VALID:
            warnungen.append('%s/%s: unbekannte Entscheidung %r ignoriert' % (ws.title, code, ent))
            ent = ''
        if art == 'Gefährdung':
            sammel = ent or None
            if sammel and sammel != 'Freigeben':
                warnungen.append('%s/%s: Sammelentscheidung %r – nur „Freigeben" '
                                 'wirkt für den ganzen Block, ignoriert' % (ws.title, code, sammel))
                sammel = None
            continue
        if art != 'Regel' or code not in offen:
            if ent and code not in offen:
                warnungen.append('%s/%s: ist nicht (mehr) offen – ignoriert' % (ws.title, code))
            continue
        wahl = ent or sammel
        if not wahl:
            continue
        if wahl in ('Ändern', 'Streichen') and not korr:
            warnungen.append('%s/%s: „%s" ohne Korrektur' % (ws.title, code, wahl))
        neu[code] = (wahl, korr if ent else '', datum, fingerabdruck(R[code], seed))

# Bestehende Datei lesen und zusammenführen
pfad = os.path.join(HERE, 'mf_content', 'regelfreigabe.py')
ns = {}
exec(open(pfad, encoding='utf-8').read(), ns)
alt_datum = ns.get('DATUM', '')
alles = {}
for code, v in ns.get('FREIGABE', {}).items():
    v = tuple(v)
    if len(v) == 2:          # Altformat: Datum aus DATUM, ohne Fingerabdruck
        v = (v[0], v[1], alt_datum, '')
    alles[code] = v
ersetzt = [c for c in neu if c in alles]
alles.update(neu)

zaehl = Counter(v[0] for v in neu.values())
print('Datei:', os.path.basename(xlsx))
print('neue Entscheidungen: %d (%s), davon %d ersetzen ältere'
      % (len(neu), ', '.join('%s %d' % kv for kv in sorted(zaehl.items())) or '–', len(ersetzt)))
print('offen danach (ohne Freigabe): %d von %d'
      % (len([c for c in offen if alles.get(c, ('',))[0] != 'Freigeben']), len(offen)))
for w in warnungen:
    print('WARNUNG', w)
aendern = [(c, v[1]) for c, v in neu.items() if v[0] != 'Freigeben']
if aendern:
    print('\nIm Inhalt (mf_content/*.py) nachzuziehen:')
    for c, k in sorted(aendern):
        print('  %-12s %s: %s' % (c, neu[c][0], k or '-'))
# Maßnahmenarten (Blatt „Maßnahmenarten“: Code, Maßnahme, …, Art richtig?)
arten = {}
if 'Maßnahmenarten' in wb.sheetnames:
    for row in wb['Maßnahmenarten'].iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue
        wahl = text(row[7]).upper()
        if wahl in ('T', 'O') and wahl != text(row[3]).upper():
            arten[row[1]] = 'TECHNICAL' if wahl == 'T' else 'ORGANISATIONAL'
if arten:
    print('\nMaßnahmenarten korrigiert: %d' % len(arten))

# Entscheidungen zum externen Prüfbericht (nur berichten und ablegen)
pb_ent = []
if 'Prüfbericht' in wb.sheetnames:
    for row in wb['Prüfbericht'].iter_rows(min_row=2, values_only=True):
        if row and row[0] and len(row) > 8 and (text(row[7]) or text(row[8])):
            pb_ent.append({'id': row[0], 'entscheidung': text(row[7]), 'kommentar': text(row[8])})
if pb_ent:
    print('\nEntscheidungen zum Prüfbericht:')
    for e in pb_ent:
        print('  %-4s %-11s %s' % (e['id'], e['entscheidung'] or '-', e['kommentar']))

if probe:
    print('\n--probe: nichts geschrieben.')
    sys.exit(0)

if arten:
    mp = os.path.join(HERE, 'mf_content', 'massnahmenart.py')
    alt = {}
    if os.path.exists(mp):
        ns2 = {}
        exec(open(mp, encoding='utf-8').read(), ns2)
        alt = ns2.get('MASSNAHMENART', {})
    alt.update(arten)
    kopfzeilen = ['# -*- coding: utf-8 -*-',
                  '# Von Hand bestätigte Maßnahmenarten (Blatt „Maßnahmenarten“ der',
                  '# Regelprüfung, übernommen mit apply_mf_regelpruefung.py). Titel -> Art.',
                  '', 'MASSNAHMENART = {']
    with open(mp, 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(kopfzeilen) + '\n')
        for k in sorted(alt):
            fh.write('    %r: %r,\n' % (k, alt[k]))
        fh.write('}\n')
    print('geschrieben:', mp)
if pb_ent:
    jp = os.path.join(HERE, 'pruefbericht_entscheidungen.json')
    with open(jp, 'w', encoding='utf-8') as fh:
        json.dump({'datum': datum, 'entscheidungen': pb_ent}, fh, ensure_ascii=False, indent=1)
    print('geschrieben:', jp)

src = open(pfad, encoding='utf-8').read()
m = re.match(r'(\s*#[^\n]*\n)*\s*"""(.*?)"""', src, re.S)
doku = m.group(0).rstrip() if m else '# -*- coding: utf-8 -*-'


def schluessel(c):
    return [int(x) if x.isdigit() else x for x in re.split(r'(\d+)', c)]


zeilen = [doku, '', '# Stand der letzten Übernahme (nur Info – jeder Eintrag trägt sein Datum)',
          'DATUM = %r' % datum, '',
          '# Code -> (Entscheidung, Korrektur, Datum, Fingerabdruck)',
          'FREIGABE = {']
for c in sorted(alles, key=schluessel):
    zeilen.append('    %r: %r,' % (c, alles[c]))
zeilen.append('}')
open(pfad, 'w', encoding='utf-8').write('\n'.join(zeilen) + '\n')
print('\ngeschrieben:', pfad, '| Einträge gesamt:', len(alles))
print('Weiter: python3 gen_mf_catalog.py && python3 gen_en8180_catalog.py')
