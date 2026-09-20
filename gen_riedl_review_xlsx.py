# -*- coding: utf-8 -*-
"""Gegenlesungs-Excel für die GBU-Variante „Riedl":
  python3 gen_riedl_review_xlsx.py            -> GBU_Riedl_Zuordnung.xlsx

Eingaben: riedl_map.json, norm_riedl_mf.json, norm_riedl_cyber.json,
          norm_81_20_mf.json, norm_cyber_mf.json (für „Nicht enthalten")
Blätter : Lesehinweise · Blaetter (je Zeile und Gefährdung: Karten/Fragen,
          TRBS-Anhang-1-Punkte, Excel-Istzustände, Prüfspalte) · Cyber ·
          Nicht enthalten (mit Prüfspalte „aufnehmen?") · Klaerungen · Umfang
"""
import json, os, subprocess, sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import riedl_content as rc  # noqa: E402  (Excel-Istzustände nur für die Gegenlesung)

EXCEL = {(b['key'], z['nr']): z.get('excel', [])
         for b in rc.BLAETTER + [rc.CYBER_BLATT] for z in b['zeilen']}


def lade(name):
    return json.load(open(os.path.join(HERE, name), encoding='utf-8'))


MAP = lade('riedl_map.json')
RG = lade('norm_riedl_mf.json')
RC = lade('norm_riedl_cyber.json')
MF = lade('norm_81_20_mf.json')
CY = lade('norm_cyber_mf.json')

FONT = 'Arial'
HEAD = PatternFill('solid', fgColor='1F3A5F')
INPUT = PatternFill('solid', fgColor='FFF2CC')
TEIL = PatternFill('solid', fgColor='FCE4E4')
STRIPE = PatternFill('solid', fgColor='F3F5F8')
BLATT = PatternFill('solid', fgColor='DDE7F3')
thin = Side(style='thin', color='C9D1DB')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')


def kopf(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 30


def zeile(ws, r, werte, eingabe=(), fill=None):
    for c, v in enumerate(werte, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=FONT, size=10)
        cell.alignment = WRAP
        cell.border = BORDER
        if c in eingabe:
            cell.fill = INPUT
        elif fill is not None:
            cell.fill = fill


def fragen_text(seed, hazard):
    Q = {q['code']: q for q in seed['questions']}
    G = {it['question']: g for g in seed.get('question_groups', []) for it in g['items']}
    out = []
    for hq in hazard['questions']:
        q = Q[hq['question']]
        karte = G.get(q['code'])
        praefix = 'Karte %s „%s"' % (karte['id'], karte['title']) if karte else 'Einzelfrage'
        out.append('%s %s (%s, %s)' % (q.get('ui_number', ''), q['text'], hq.get('role', ''), praefix))
    return '\n'.join(out)


def blatt_zeilen(ws, blaetter, seed, cyber=False):
    cols = ['Blatt', 'Nr.', 'Zeile (Bericht)', 'VFA-Nr.' if not cyber else 'Komponente',
            'Gefährdung', 'Titel der Gefährdung', 'Fragen / Karten', 'TRBS 3121 Anh. 1',
            'Deckung', 'Excel-Istzustände, die die Zeile ersetzt', 'Bemerkung',
            'Prüfung', 'Kommentar / Korrektur']
    widths = [16, 5, 34, 11, 10, 46, 70, 10, 10, 46, 46, 12, 34]
    kopf(ws, cols, widths)
    H = {h['code']: h for h in seed['hazards']}
    r = 2
    for b in blaetter:
        for z in b['zeilen']:
            erster = True
            for c in z['hazards']:
                h = H[c]
                werte = [b['titel'], z['nr'], z['titel'],
                         (z.get('komponente') or '-') if cyber else (z.get('vfa') or '-'),
                         c, h['title'], fragen_text(seed, h),
                         ', '.join(str(n) for n in z['trbs_anhang1']) or '-',
                         z['deckung'],
                         '\n'.join(EXCEL.get((b['key'], z['nr']), [])) if erster else '(s. o.)',
                         z.get('bemerkung', '') if erster else '',
                         None, None]
                fill = TEIL if z['deckung'] != 'voll' else (BLATT if erster else (STRIPE if r % 2 == 0 else None))
                zeile(ws, r, werte, eingabe=(12, 13), fill=fill)
                erster = False
                r += 1
    dv = DataValidation(type='list', formula1='"OK,Ändern,Streichen"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add('L2:L%d' % (r - 1))
    ws.auto_filter.ref = 'A1:M%d' % (r - 1)
    return r - 1


wb = Workbook()

# ---- Lesehinweise -----------------------------------------------------------
ws = wb.active
ws.title = 'Lesehinweise'
ws.column_dimensions['A'].width = 125
umfang_gbu = '%d Fragen, %d Gefährdungen, %d Regeln, %d Karten' % (
    len(RG['questions']), len(RG['hazards']), len(RG['rules']), len(RG.get('question_groups', [])))
umfang_cy = '%d Fragen, %d Gefährdungen, %d Regeln' % (
    len(RC['questions']), len(RC['hazards']), len(RC['rules']))
hinweise = [
    ('GBU-Variante „Riedl" – Zuordnung der VFA-Blätter und der Cyber-Komponenten (Stand %s)' % '17.09.2026', True),
    ('Kataloge: %s (aus %s) und %s (aus %s)' % (MAP['kataloge']['gbu'], MAP['kataloge']['gbu_quelle'],
                                                 MAP['kataloge']['cyber'], MAP['kataloge']['cyber_quelle']), False),
    ('', False),
    ('Die Variante ersetzt die beiden Excel-Vorlagen (VFA „GBU Anlage leer" R 1.7.2 und „Vorlage GBU '
     'Cybersicherheit ab 04-26"). Sie ist ABGELEITET: jede Gefährdung, Frage, Regel und Maßnahme ist '
     'byteweise identisch mit dem MF- bzw. Cyber-Katalog; nur der Umfang ist kleiner. Umfang GBU: %s '
     '(MF: %d/%d/%d/%d). Umfang Cyber: %s (CY: %d/%d/%d).' % (
         umfang_gbu, len(MF['questions']), len(MF['hazards']), len(MF['rules']), len(MF.get('question_groups', [])),
         umfang_cy, len(CY['questions']), len(CY['hazards']), len(CY['rules'])), False),
    ('', False),
    ('Umfangsregel: Eine MF-Gefährdung ist drin, wenn (a) eine Zeile oder ein Istzustand der VFA-Vorlage sie '
     'abfragt, (b) sie einen der 22 Punkte des TRBS-3121-Anhangs 1 trägt (das VFA-Blatt „Bezug zum Stand der '
     'Technik" führt genau diese) oder (c) sie eine Betreiberpflicht ist, die der Block Z / das Deckblatt '
     'voraussetzt. Alles andere steht im Blatt „Nicht enthalten" mit Begründung.', False),
    ('', False),
    ('Bewertung: Engine-Stufe als Ampel (Kein Risiko/Niedrig = Grün, Mittel = Gelb, Hoch = Rot). Eine Zeile mit '
     'mehreren Gefährdungen zeigt die schlechteste Farbe; nicht zutreffende Gefährdungen zählen nicht; eine '
     'unvollständige Gefährdung hält die Zeile „offen". Blatt- und Gesamtampel wie die Deckblatt-Formel der '
     'Vorlage (alles Grün = Grün, ein Rot = Rot, sonst Gelb).', False),
    ('', False),
    ('Blatt „Blaetter": je Zeile des Berichts und je zugeordneter Gefährdung eine Zeile. Spalte „Fragen / Karten" '
     'zeigt, was der Techniker dazu beantwortet (Karte = ein Tipp „Keine weiteren Auffälligkeiten" oder ein '
     'Ankreuzfeld). Spalte „Excel-Istzustände" zeigt die Auswahlwerte der Vorlage, die die Zeile ersetzt (nur '
     'Vergleich, nicht übernommen). Rosa = nur teilweise abgedeckt (Bemerkung lesen).', False),
    ('Blatt „Cyber": dieselbe Struktur für die sieben Riedl-Komponenten plus Zugang, Netz, Organisation. Die '
     'Spalte „Erläuterung" ist der Standardtext, der bei Grün im Bericht erscheint (eigene Formulierung nach '
     'dem Blatt „Erläuterungen für Standard").', False),
    ('Blatt „Nicht enthalten": alle MF-/CY-Gefährdungen, die bewusst draußen bleiben – mit Begründung und der '
     'Möglichkeit, sie in eine Zeile zu ziehen (Spalte „aufnehmen in Zeile").', False),
    ('Blatt „Klaerungen": offene Entscheidungen K-R01 … K-R08 mit Vorschlag und Alternative.', False),
    ('Blatt „Umfang": Zahlen je Erhebungsbereich und Simulation, was ein Prüfer je Anlagenprofil sieht.', False),
    ('', False),
    ('Gelbe Spalten ausfüllen: „Prüfung" = OK / Ändern / Streichen bzw. „Entscheidung", Kommentar frei. '
     'Rücklauf: Datei zurückgeben, Änderungen werden in riedl_content.py eingepflegt und der Generator '
     'neu ausgeführt (python3 gen_riedl_catalog.py).', False),
]
for i, (t, b) in enumerate(hinweise, 1):
    cell = ws.cell(row=i, column=1, value=t)
    cell.font = Font(name=FONT, bold=b, size=11 if b else 10)
    cell.alignment = WRAP

# ---- Blaetter ----------------------------------------------------------------
n_blatt = blatt_zeilen(wb.create_sheet('Blaetter'), MAP['blaetter'], RG)

# ---- Cyber --------------------------------------------------------------------
wsc = wb.create_sheet('Cyber')
n_cyber = blatt_zeilen(wsc, [MAP['cyber']], RC, cyber=True)
# Erläuterung als Zusatzspalte
wsc.cell(row=1, column=14, value='Erläuterung (Standardtext bei Grün)').font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
wsc.cell(row=1, column=14).fill = HEAD
wsc.cell(row=1, column=14).border = BORDER
wsc.column_dimensions['N'].width = 70
r = 2
for z in MAP['cyber']['zeilen']:
    for i, _ in enumerate(z['hazards']):
        cell = wsc.cell(row=r, column=14, value=z.get('erlaeuterung', '') if i == 0 else '')
        cell.font = Font(name=FONT, size=10)
        cell.alignment = WRAP
        cell.border = BORDER
        r += 1
wsc.auto_filter.ref = 'A1:N%d' % (r - 1)

# ---- Nicht enthalten ----------------------------------------------------------
wsn = wb.create_sheet('Nicht enthalten')
kopf(wsn, ['Gefährdung', 'Titel', 'Erhebungsbereich', 'Fragen', 'Begründung', 'Entscheidung',
           'aufnehmen in Zeile (Blatt/Nr.)', 'Kommentar'],
     [11, 52, 30, 7, 60, 14, 22, 34])
Hq = {h['code']: h for h in MF['hazards'] + CY['hazards']}
r = 2
for code, grund in MAP['nicht_enthalten'].items():
    h = Hq[code]
    zeile(wsn, r, [code, h['title'], h.get('ui_group', ''), len(h['questions']), grund, None, None, None],
          eingabe=(6, 7, 8), fill=STRIPE if r % 2 == 0 else None)
    r += 1
dv = DataValidation(type='list', formula1='"bleibt draußen,aufnehmen"', allow_blank=True)
wsn.add_data_validation(dv)
dv.add('F2:F%d' % (r - 1))
wsn.auto_filter.ref = 'A1:H%d' % (r - 1)

# ---- Klaerungen ---------------------------------------------------------------
wsk = wb.create_sheet('Klaerungen')
kopf(wsk, ['ID', 'Thema', 'Vorschlag', 'Alternative', 'Entscheidung', 'Kommentar'],
     [8, 30, 70, 60, 14, 40])
r = 2
for k in MAP['klaerungen']:
    zeile(wsk, r, [k['id'], k['thema'], k['vorschlag'], k['alternative'], None, None],
          eingabe=(5, 6), fill=STRIPE if r % 2 == 0 else None)
    r += 1
dv = DataValidation(type='list', formula1='"Vorschlag,Alternative,Anders (Kommentar)"', allow_blank=True)
wsk.add_data_validation(dv)
dv.add('E2:E%d' % (r - 1))

# ---- Umfang ----------------------------------------------------------------------
wsu = wb.create_sheet('Umfang')
wsu.column_dimensions['A'].width = 48
wsu.column_dimensions['B'].width = 18
wsu.column_dimensions['C'].width = 18
wsu.column_dimensions['D'].width = 18
r = 1
def kopfzeile(text):
    global r
    c = wsu.cell(row=r, column=1, value=text)
    c.font = Font(name=FONT, bold=True, size=11)
    r += 1
def wert(a, b='', c='', d=''):
    global r
    for i, v in enumerate((a, b, c, d), 1):
        cell = wsu.cell(row=r, column=i, value=v)
        cell.font = Font(name=FONT, size=10)
        cell.alignment = WRAP
    r += 1

kopfzeile('Umfang je Erhebungsbereich (Fragen) – Riedl gegenüber MF')
wert('Erhebungsbereich', 'Riedl', 'MF (voll)')
kat_r = Counter(q['category'] for q in RG['questions'])
kat_m = Counter(q['category'] for q in MF['questions'])
for k in MF.get('category_order', sorted(kat_m)):
    wert(k, kat_r.get(k, 0), kat_m.get(k, 0))
wert('Summe', len(RG['questions']), len(MF['questions']))
wert('Gefährdungen', len(RG['hazards']), len(MF['hazards']))
wert('Regeln', len(RG['rules']), len(MF['rules']))
wert('Karten', len(RG.get('question_groups', [])), len(MF.get('question_groups', [])))
wert('Nachweise (ZÜS-Vorbelegung)', len(RG.get('nachweise', [])), len(MF.get('nachweise', [])))
wert('Begründete Annahmen (Baujahr)', len(RG.get('assumptions', [])), len(MF.get('assumptions', [])))
r += 1
kopfzeile('Cyber – Riedl gegenüber CY')
wert('Erhebungsbereich', 'Riedl', 'CY (voll)')
kat_rc = Counter(q['category'] for q in RC['questions'])
kat_c = Counter(q['category'] for q in CY['questions'])
for k in sorted(kat_c):
    wert(k, kat_rc.get(k, 0), kat_c.get(k, 0))
wert('Summe', len(RC['questions']), len(CY['questions']))
wert('Gefährdungen', len(RC['hazards']), len(CY['hazards']))
wert('Regeln', len(RC['rules']), len(CY['rules']))
r += 1
kopfzeile('Simulation: Was sieht der Prüfer je Anlagenprofil? (sim_karten.py)')
wert('Profil', 'Karten vor Ort offen (Riedl)', 'Karten vor Ort offen (MF)', 'Fragen sichtbar (Riedl / MF)')


def sim(seed):
    out = subprocess.run([sys.executable, os.path.join(HERE, 'sim_karten.py'), seed],
                         capture_output=True, text=True, cwd=HERE).stdout
    ergebnis = {}
    name = None
    for line in out.splitlines():
        if line.startswith('P'):
            name = line.strip()
        elif 'Fragen sichtbar' in line and name:
            teile = dict(t.strip().split(': ', 1) for t in line.split('|') if ': ' in t)
            ergebnis[name] = (teile.get('Fragen sichtbar'), teile.get('vor Ort'))
    return ergebnis


sr = sim('norm_riedl_mf.json')
sm = sim('norm_81_20_mf.json')
for name in sr:
    wert(name, sr[name][1], sm.get(name, ('', ''))[1], '%s / %s' % (sr[name][0], sm.get(name, ('', ''))[0]))
wert('Zum Vergleich: Excel-Vorlage', '≈ 28 Auswahlzeilen GBU + 7 Cyber-Komponenten')

ziel = os.path.join(HERE, 'GBU_Riedl_Zuordnung.xlsx')
wb.save(ziel)
print('geschrieben: %s (Blaetter %d Zeilen, Cyber %d Zeilen, Nicht enthalten %d, Klärungen %d)'
      % (os.path.basename(ziel), n_blatt - 1, n_cyber - 1, len(MAP['nicht_enthalten']), len(MAP['klaerungen'])))
